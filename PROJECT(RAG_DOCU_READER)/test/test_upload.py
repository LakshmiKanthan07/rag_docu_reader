"""
test_upload.py — File ingestion pipeline
=========================================
Answers: "Can data enter the system correctly?"

Covers:
  - Valid file types are accepted
  - Unsupported / oversized / empty files are rejected gracefully
  - Response shape is correct (session_id, total_chunks, files[])
  - Multiple files are merged into one session
  - Per-file status is reported individually
  - Processing pipeline: extract → chunk → embed → store
  - Session is queryable after upload
"""

import io
import os
import shutil
import tempfile
import pytest
from unittest.mock import patch, MagicMock

from rag_docu_reader import main


# ── Helpers ────────────────────────────────────────────────────────────────────

def _upload(client, files: list[tuple[str, bytes, str]]):
    """Helper: POST /upload with one or more (filename, bytes, mime) tuples."""
    return client.post(
        "/upload",
        files=[("files", (name, io.BytesIO(data), mime)) for name, data, mime in files],
    )


# ── Response shape ─────────────────────────────────────────────────────────────

class TestUploadResponseShape:
    def test_returns_200(self, client, txt_file):
        resp = _upload(client, [txt_file])
        assert resp.status_code == 200

    def test_response_has_session_id(self, client, txt_file):
        data = _upload(client, [txt_file]).json()
        assert "session_id" in data
        assert len(data["session_id"]) > 0

    def test_response_has_total_chunks(self, client, txt_file):
        data = _upload(client, [txt_file]).json()
        assert "total_chunks" in data
        assert isinstance(data["total_chunks"], int)
        assert data["total_chunks"] > 0

    def test_response_has_files_list(self, client, txt_file):
        data = _upload(client, [txt_file]).json()
        assert "files" in data
        assert isinstance(data["files"], list)
        assert len(data["files"]) == 1

    def test_response_has_message(self, client, txt_file):
        data = _upload(client, [txt_file]).json()
        assert "message" in data

    def test_per_file_status_ok(self, client, txt_file):
        files_info = _upload(client, [txt_file]).json()["files"]
        assert files_info[0]["status"] == "ok"
        assert "chunks" in files_info[0]

    def test_per_file_filename_matches(self, client, txt_file):
        filename, _, _ = txt_file
        files_info = _upload(client, [txt_file]).json()["files"]
        assert files_info[0]["file"] == filename


# ── Supported file types ───────────────────────────────────────────────────────

class TestSupportedFileTypes:
    def test_txt_file_accepted(self, client, txt_file):
        assert _upload(client, [txt_file]).status_code == 200

    def test_md_file_accepted(self, client, md_file):
        assert _upload(client, [md_file]).status_code == 200

    def test_csv_file_accepted(self, client, csv_file):
        assert _upload(client, [csv_file]).status_code == 200

    def test_json_file_accepted(self, client, json_file):
        assert _upload(client, [json_file]).status_code == 200

    def test_py_file_accepted(self, client):
        py = ("script.py", b"def hello():\n    return 'world'\n", "text/plain")
        assert _upload(client, [py]).status_code == 200

    def test_xml_file_accepted(self, client):
        xml = ("data.xml", b"<root><item>hello</item></root>", "application/xml")
        assert _upload(client, [xml]).status_code == 200

    def test_yaml_file_accepted(self, client):
        yml = ("config.yaml", b"key: value\nlist:\n  - a\n  - b\n", "text/yaml")
        assert _upload(client, [yml]).status_code == 200


# ── Unsupported / bad input ────────────────────────────────────────────────────

class TestUnsupportedFiles:
    def test_exe_file_returns_4xx(self, client, unsupported_file):
        resp = _upload(client, [unsupported_file])
        # Could be 422 (skipped, no usable content) or 415 (rejected early)
        assert resp.status_code in (415, 422)

    def test_unsupported_file_has_skipped_status(self, client, unsupported_file):
        resp = _upload(client, [unsupported_file])
        data = resp.json()
        # The detail may be a dict with "files" key when 422
        if resp.status_code == 422 and isinstance(data.get("detail"), dict):
            files_info = data["detail"]["files"]
        else:
            files_info = data.get("files", [])
        if files_info:
            assert files_info[0]["status"] in ("skipped", "error")

    def test_zip_file_rejected(self, client):
        zf = ("archive.zip", b"PK\x03\x04", "application/zip")
        resp = _upload(client, [zf])
        assert resp.status_code in (415, 422)


# ── Empty and oversized files ──────────────────────────────────────────────────

class TestEdgeCaseFiles:
    def test_empty_file_returns_422(self, client, empty_file):
        resp = _upload(client, [empty_file])
        assert resp.status_code == 422

    def test_empty_file_error_message_present(self, client, empty_file):
        resp = _upload(client, [empty_file])
        detail = resp.json().get("detail", "")
        # Should mention "no usable content" or similar
        assert detail != ""

    def test_large_file_returns_4xx(self, client, large_file):
        resp = _upload(client, [large_file])
        # Expect 400/413/422 — must not be 200 or 500
        assert resp.status_code in (400, 413, 422)
        assert resp.status_code != 500

    def test_no_files_returns_400(self, client):
        resp = client.post("/upload")
        assert resp.status_code in (400, 422)


# ── Multi-file upload ──────────────────────────────────────────────────────────

class TestMultiFileUpload:
    def test_two_files_return_one_session(self, client, txt_file, csv_file):
        data = _upload(client, [txt_file, csv_file]).json()
        assert "session_id" in data
        assert len(data["session_id"]) > 0

    def test_two_files_chunks_are_combined(self, client, txt_file, csv_file):
        resp_single = _upload(client, [txt_file]).json()["total_chunks"]
        resp_double = _upload(client, [txt_file, csv_file]).json()["total_chunks"]
        assert resp_double >= resp_single

    def test_per_file_results_length_matches_files_sent(self, client, txt_file, csv_file):
        data = _upload(client, [txt_file, csv_file]).json()
        assert len(data["files"]) == 2

    def test_mixed_valid_invalid_partial_success(self, client, txt_file, unsupported_file):
        """Valid + invalid: session created from valid file; invalid is skipped."""
        resp = _upload(client, [txt_file, unsupported_file])
        # Valid file should have saved the session
        assert resp.status_code == 200
        data = resp.json()
        statuses = [f["status"] for f in data["files"]]
        assert "ok" in statuses
        assert "skipped" in statuses or "error" in statuses


# ── Session creation and isolation ────────────────────────────────────────────

class TestSessionCreation:
    def test_each_upload_creates_unique_session(self, client, txt_file):
        id1 = _upload(client, [txt_file]).json()["session_id"]
        id2 = _upload(client, [txt_file]).json()["session_id"]
        assert id1 != id2

    def test_session_is_queryable_after_upload(self, client, seeded_session):
        """A freshly created session must respond to /ask without 404."""
        resp = client.post(
            "/ask",
            json={"session_id": seeded_session, "question": "What is this about?"},
        )
        assert resp.status_code == 200

    def test_session_deleted_after_delete_call(self, client, seeded_session):
        del_resp = client.delete(f"/session/{seeded_session}")
        assert del_resp.status_code == 200
        # Now /ask should fail with 404
        ask_resp = client.post(
            "/ask",
            json={"session_id": seeded_session, "question": "hi"},
        )
        assert ask_resp.status_code == 404


# ── Chunk metadata ─────────────────────────────────────────────────────────────

class TestChunkMetadata:
    def test_chunks_count_is_positive(self, client, txt_file):
        data = _upload(client, [txt_file]).json()
        assert data["total_chunks"] >= 1

    def test_large_text_produces_multiple_chunks(self, client):
        """A file with 10× the chunk size should be split."""
        big_text = ("big.txt", ("word " * (main.CHUNK_SIZE * 10)).encode(), "text/plain")
        data = _upload(client, [big_text]).json()
        assert data["total_chunks"] > 1


# ── Docx / xlsx (stub loaders) ─────────────────────────────────────────────────

class TestDocxXlsxLoaders:
    """
    Real docx/xlsx bytes are generated with python-docx / openpyxl.
    If those aren't installed in the test env, tests are skipped gracefully.
    """

    def test_docx_upload(self, client, tmp_path):
        pytest.importorskip("docx", reason="python-docx not installed")
        import docx as _docx
        doc = _docx.Document()
        doc.add_paragraph("This is a test Word document about neural networks.")
        path = tmp_path / "test.docx"
        doc.save(str(path))
        with open(str(path), "rb") as f:
            content = f.read()
        resp = _upload(client, [("test.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")])
        assert resp.status_code == 200

    def test_xlsx_upload(self, client, tmp_path):
        pytest.importorskip("openpyxl", reason="openpyxl not installed")
        import openpyxl as _xl
        wb = _xl.Workbook()
        ws = wb.active
        ws.append(["Name", "Score"])
        ws.append(["Alice", 95])
        ws.append(["Bob", 87])
        path = tmp_path / "test.xlsx"
        wb.save(str(path))
        with open(str(path), "rb") as f:
            content = f.read()
        resp = _upload(client, [("test.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")])
        assert resp.status_code == 200